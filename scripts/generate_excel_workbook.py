import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open("/Users/ridwanilyas/Documents/Unjani/Statistika/bank_soal_300.json", "r", encoding="utf-8") as f:
    data = json.load(f)

all_pg = data["pilihan_ganda"]
all_is = data["isian_singkat"]
all_es = data["essay"]

wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Styling definitions
header_fill_blue = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
header_fill_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
header_fill_green = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_fill_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
header_fill_indigo = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
title_font = Font(name="Calibri", size=14, bold=True, color="1A365D")
sub_font = Font(name="Calibri", size=11, italic=True, color="4A5568")
bold_font = Font(name="Calibri", size=10, bold=True)
regular_font = Font(name="Calibri", size=10)

thin_border = Border(
    left=Side(style='thin', color='CBD5E0'),
    right=Side(style='thin', color='CBD5E0'),
    top=Side(style='thin', color='CBD5E0'),
    bottom=Side(style='thin', color='CBD5E0')
)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

def style_sheet(ws, header_fill):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws.row_dimensions[1].height = 28
    
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 24
        is_even = (row % 2 == 0)
        for col in range(1, ws.max_column + 1):
            c = ws.cell(row=row, column=col)
            c.font = regular_font
            c.border = thin_border
            if is_even:
                c.fill = zebra_fill
            if col in [1, 2, 4]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ----------------------------------------------------
# 1. Sheet Panduan
# ----------------------------------------------------
ws_info = wb.create_sheet(title="Panduan & Distribusi")
ws_info.views.sheetView[0].showGridLines = True
ws_info.cell(row=1, column=1, value="BANK SOAL 300 STATISTIKA KOMPUTASI (OBE TEKNIK INFORMATIKA)").font = title_font
ws_info.cell(row=2, column=1, value="Mata Kuliah: Statistika Komputasi | Penyusun: Dr. Ridwan Ilyas, S.Kom., M.T. | UNJANI").font = sub_font
ws_info.cell(row=4, column=1, value="Struktur Bank Soal:").font = bold_font
ws_info.cell(row=5, column=1, value="1. Bagian I: 100 Soal Pilihan Ganda (PG001 - PG100)").font = regular_font
ws_info.cell(row=6, column=1, value="2. Bagian II: 100 Soal Isian Singkat & Kasus Komputasi (IS001 - IS100)").font = regular_font
ws_info.cell(row=7, column=1, value="3. Bagian III: 100 Soal Essay & Studi Kasus Analisis (ES001 - ES100)").font = regular_font
ws_info.cell(row=8, column=1, value="4. Bagian IV: Kunci Jawaban & Pembahasan Lengkap (Terpisah pada Sheet ke-5)").font = bold_font

ws_info.cell(row=10, column=1, value="Tabel Distribusi 25 Modul (4 PG + 4 Isian + 4 Essay = 12 Soal per Modul):").font = bold_font

headers_dist = ["No", "Buku", "Modul ID", "Nama Modul", "Pilihan Ganda", "Isian Singkat", "Essay", "Total Soal"]
for c_idx, h in enumerate(headers_dist, 1):
    cell = ws_info.cell(row=11, column=c_idx, value=h)
    cell.fill = header_fill_blue
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

row_idx = 12
modules_summary = []
for m_id in range(1, 26):
    pg_sample = [q for q in all_pg if q["modul_id"] == m_id][0]
    ws_info.cell(row=row_idx, column=1, value=m_id)
    ws_info.cell(row=row_idx, column=2, value=pg_sample["buku"])
    ws_info.cell(row=row_idx, column=3, value=f"Modul {m_id:02d}")
    ws_info.cell(row=row_idx, column=4, value=pg_sample["modul"])
    ws_info.cell(row=row_idx, column=5, value=4)
    ws_info.cell(row=row_idx, column=6, value=4)
    ws_info.cell(row=row_idx, column=7, value=4)
    ws_info.cell(row=row_idx, column=8, value=12)
    for c in range(1, 9):
        ws_info.cell(row=row_idx, column=c).border = thin_border
        if c in [1, 3, 5, 6, 7, 8]:
            ws_info.cell(row=row_idx, column=c).alignment = Alignment(horizontal="center")
    row_idx += 1

# Total Row
ws_info.cell(row=row_idx, column=1, value="Total")
ws_info.cell(row=row_idx, column=4, value="TOTAL KESELURUHAN BANK SOAL")
ws_info.cell(row=row_idx, column=5, value=100)
ws_info.cell(row=row_idx, column=6, value=100)
ws_info.cell(row=row_idx, column=7, value=100)
ws_info.cell(row=row_idx, column=8, value=300)
for c in range(1, 9):
    ws_info.cell(row=row_idx, column=c).font = bold_font
    ws_info.cell(row=row_idx, column=c).fill = zebra_fill
    ws_info.cell(row=row_idx, column=c).border = thin_border

# ----------------------------------------------------
# 2. Sheet 100 Pilihan Ganda (Soal Only)
# ----------------------------------------------------
ws_pg = wb.create_sheet(title="100 Pilihan Ganda")
ws_pg.views.sheetView[0].showGridLines = True
headers_pg = ["No", "Kode Soal", "Buku", "Modul ID", "Nama Modul", "Pertanyaan / Kasus", "Pilihan A", "Pilihan B", "Pilihan C", "Pilihan D", "Pilihan E"]
for col, h in enumerate(headers_pg, 1):
    ws_pg.cell(row=1, column=col, value=h)

for idx, q in enumerate(all_pg, 1):
    row = idx + 1
    ws_pg.cell(row=row, column=1, value=idx)
    ws_pg.cell(row=row, column=2, value=q["id"])
    ws_pg.cell(row=row, column=3, value=q["buku"])
    ws_pg.cell(row=row, column=4, value=f"Modul {q['modul_id']:02d}")
    ws_pg.cell(row=row, column=5, value=q["modul"])
    ws_pg.cell(row=row, column=6, value=q["soal"])
    ws_pg.cell(row=row, column=7, value=q["opsi"].get("A", ""))
    ws_pg.cell(row=row, column=8, value=q["opsi"].get("B", ""))
    ws_pg.cell(row=row, column=9, value=q["opsi"].get("C", ""))
    ws_pg.cell(row=row, column=10, value=q["opsi"].get("D", ""))
    ws_pg.cell(row=row, column=11, value=q["opsi"].get("E", ""))

style_sheet(ws_pg, header_fill_blue)

# ----------------------------------------------------
# 3. Sheet 100 Isian Singkat (Soal Only)
# ----------------------------------------------------
ws_is = wb.create_sheet(title="100 Isian Singkat")
ws_is.views.sheetView[0].showGridLines = True
headers_is = ["No", "Kode Soal", "Buku", "Modul ID", "Nama Modul", "Pertanyaan / Kasus Komputasi Singkat"]
for col, h in enumerate(headers_is, 1):
    ws_is.cell(row=1, column=col, value=h)

for idx, q in enumerate(all_is, 1):
    row = idx + 1
    ws_is.cell(row=row, column=1, value=idx)
    ws_is.cell(row=row, column=2, value=q["id"])
    ws_is.cell(row=row, column=3, value=q["buku"])
    ws_is.cell(row=row, column=4, value=f"Modul {q['modul_id']:02d}")
    ws_is.cell(row=row, column=5, value=q["modul"])
    ws_is.cell(row=row, column=6, value=q["soal"])

style_sheet(ws_is, header_fill_teal)

# ----------------------------------------------------
# 4. Sheet 100 Soal Essay (Soal Only)
# ----------------------------------------------------
ws_es = wb.create_sheet(title="100 Soal Essay")
ws_es.views.sheetView[0].showGridLines = True
headers_es = ["No", "Kode Soal", "Buku", "Modul ID", "Nama Modul", "Kasus & Instruksi Soal Essay Analitis"]
for col, h in enumerate(headers_es, 1):
    ws_es.cell(row=1, column=col, value=h)

for idx, q in enumerate(all_es, 1):
    row = idx + 1
    ws_es.cell(row=row, column=1, value=idx)
    ws_es.cell(row=row, column=2, value=q["id"])
    ws_es.cell(row=row, column=3, value=q["buku"])
    ws_es.cell(row=row, column=4, value=f"Modul {q['modul_id']:02d}")
    ws_es.cell(row=row, column=5, value=q["modul"])
    ws_es.cell(row=row, column=6, value=q["soal"])

style_sheet(ws_es, header_fill_indigo)

# ----------------------------------------------------
# 5. Sheet Kunci Jawaban & Pembahasan (Terpisah 300 Soal)
# ----------------------------------------------------
ws_ans = wb.create_sheet(title="Kunci Jawaban & Pembahasan")
ws_ans.views.sheetView[0].showGridLines = True
headers_ans = ["No", "Kode Soal", "Tipe Soal", "Buku", "Modul ID", "Nama Modul", "Kunci Jawaban / Solusi", "Pembahasan / Rubrik Penilaian"]
for col, h in enumerate(headers_ans, 1):
    ws_ans.cell(row=1, column=col, value=h)

row_ans = 2
counter_all = 1

# PG answers
for q in all_pg:
    ws_ans.cell(row=row_ans, column=1, value=counter_all)
    ws_ans.cell(row=row_ans, column=2, value=q["id"])
    ws_ans.cell(row=row_ans, column=3, value="Pilihan Ganda")
    ws_ans.cell(row=row_ans, column=4, value=q["buku"])
    ws_ans.cell(row=row_ans, column=5, value=f"Modul {q['modul_id']:02d}")
    ws_ans.cell(row=row_ans, column=6, value=q["modul"])
    ws_ans.cell(row=row_ans, column=7, value=f"{q['kunci']}. {q['opsi'].get(q['kunci'], '')}")
    ws_ans.cell(row=row_ans, column=8, value=q["pembahasan"])
    row_ans += 1
    counter_all += 1

# IS answers
for q in all_is:
    ws_ans.cell(row=row_ans, column=1, value=counter_all)
    ws_ans.cell(row=row_ans, column=2, value=q["id"])
    ws_ans.cell(row=row_ans, column=3, value="Isian Singkat")
    ws_ans.cell(row=row_ans, column=4, value=q["buku"])
    ws_ans.cell(row=row_ans, column=5, value=f"Modul {q['modul_id']:02d}")
    ws_ans.cell(row=row_ans, column=6, value=q["modul"])
    ws_ans.cell(row=row_ans, column=7, value=q["kunci"])
    ws_ans.cell(row=row_ans, column=8, value=q["pembahasan"])
    row_ans += 1
    counter_all += 1

# ES answers
for q in all_es:
    ws_ans.cell(row=row_ans, column=1, value=counter_all)
    ws_ans.cell(row=row_ans, column=2, value=q["id"])
    ws_ans.cell(row=row_ans, column=3, value="Essay")
    ws_ans.cell(row=row_ans, column=4, value=q["buku"])
    ws_ans.cell(row=row_ans, column=5, value=f"Modul {q['modul_id']:02d}")
    ws_ans.cell(row=row_ans, column=6, value=q["modul"])
    ws_ans.cell(row=row_ans, column=7, value=q["kunci"])
    ws_ans.cell(row=row_ans, column=8, value=q["rubrik"])
    row_ans += 1
    counter_all += 1

style_sheet(ws_ans, header_fill_amber)

# Set Column Widths across all sheets
col_widths = {
    "Panduan & Distribusi": {1: 8, 2: 12, 3: 14, 4: 45, 5: 16, 6: 16, 7: 16, 8: 16},
    "100 Pilihan Ganda": {1: 6, 2: 12, 3: 10, 4: 12, 5: 35, 6: 55, 7: 25, 8: 25, 9: 25, 10: 25, 11: 25},
    "100 Isian Singkat": {1: 6, 2: 12, 3: 10, 4: 12, 5: 35, 6: 75},
    "100 Soal Essay": {1: 6, 2: 12, 3: 10, 4: 12, 5: 35, 6: 85},
    "Kunci Jawaban & Pembahasan": {1: 6, 2: 12, 3: 16, 4: 10, 5: 12, 6: 35, 7: 35, 8: 75}
}

for sheet_name, widths in col_widths.items():
    ws = wb[sheet_name]
    for col_idx, width in widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

excel_path = "/Users/ridwanilyas/Documents/Unjani/Statistika/latihan_soal_300_statistika_komputasi.xlsx"
wb.save(excel_path)
print(f"Successfully generated: {excel_path}")
